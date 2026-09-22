# -*- coding: utf-8 -*-
"""
lector_excel.py — Lectura robusta de archivos Excel.

Responsabilidad única: abrir libros y hojas de Excel con mensajes de error
claros para el usuario final (archivo inexistente, hoja inexistente, archivo
dañado o abierto en Excel).
"""

from __future__ import annotations

from pathlib import Path
from collections import Counter

import pandas as pd
from openpyxl import load_workbook

from logs import obtener_logger


class ErrorLectura(Exception):
    """Error de lectura de Excel con mensaje apto para mostrar al usuario."""


def listar_hojas(ruta: Path) -> list[str]:
    """Devuelve los nombres de hoja de un libro sin cargar los datos."""
    try:
        wb = load_workbook(ruta, read_only=True)
        hojas = list(wb.sheetnames)
        wb.close()
        return hojas
    except Exception as exc:
        raise ErrorLectura(
            f"No se pudo abrir '{ruta.name}'. Verifique que no esté abierto en "
            f"Excel y que no esté dañado.\nDetalle: {exc}"
        ) from exc


def leer_columnas(ruta: Path, hoja: str) -> list[str]:
    """Lee solo los encabezados de una hoja de Excel."""
    try:
        df = pd.read_excel(ruta, sheet_name=hoja, nrows=0, engine="openpyxl")
        return [str(c).strip() for c in df.columns]
    except Exception:
        return []


def leer_valores_columna(
    ruta: Path, hoja: str, columna: str, limite: int = 100
) -> list[str]:
    """Lee los valores distintos de una sola columna sin cargar toda la hoja."""
    try:
        wb = load_workbook(ruta, read_only=True, data_only=True)
        ws = wb[hoja]
        filas = ws.iter_rows(values_only=True)
        encabezados = [str(v).strip() if v is not None else "" for v in next(filas)]
        if columna not in encabezados:
            wb.close()
            return []

        indice = encabezados.index(columna)
        frecuencias: Counter[str] = Counter()
        for fila in filas:
            valor = fila[indice] if indice < len(fila) else None
            if valor is None:
                continue
            if isinstance(valor, float) and valor.is_integer():
                valor = int(valor)
            texto = str(valor).strip()
            if texto:
                frecuencias[texto] += 1
            if len(frecuencias) >= limite:
                break
        wb.close()
        return [valor for valor, _ in frecuencias.most_common(limite)]
    except Exception:
        return []


def leer_hoja(ruta: Path, hoja: str, descripcion: str) -> pd.DataFrame:
    """Lee una hoja completa como DataFrame, validando su existencia."""
    log = obtener_logger()

    if not ruta.exists():
        raise ErrorLectura(
            f"No se encontró el archivo de {descripcion}:\n{ruta}\n\n"
            f"Coloque el archivo dentro de la carpeta 'Entrada'."
        )

    hojas = listar_hojas(ruta)
    if hoja not in hojas:
        raise ErrorLectura(
            f"La hoja '{hoja}' no existe en '{ruta.name}'.\n"
            f"Hojas disponibles: {', '.join(hojas)}"
        )

    log.info("Leyendo %s: %s (hoja '%s')...", descripcion, ruta.name, hoja)
    try:
        df = pd.read_excel(ruta, sheet_name=hoja, engine="openpyxl")
    except Exception as exc:
        raise ErrorLectura(
            f"Error al leer la hoja '{hoja}' de '{ruta.name}'.\nDetalle: {exc}"
        ) from exc

    # Encabezados sin espacios accidentales ("Municipio " -> "Municipio").
    df.columns = [c.strip() if isinstance(c, str) else c for c in df.columns]
    log.info("  -> %s registros, %s columnas.", f"{len(df):,}", len(df.columns))
    return df

# =============================================================================
# UTILIDADES PARA LA INTERFAZ
# =============================================================================

def listar_archivos_excel(carpeta: Path) -> list[str]:
    """
    Devuelve todos los archivos Excel encontrados en una carpeta.

    Se usa para llenar los ComboBox de selección de archivos
    en la interfaz gráfica.
    """
    extensiones = ("*.xlsx", "*.xls", "*.xlsm", "*.xlsb")

    archivos = []

    for ext in extensiones:
        archivos.extend(carpeta.glob(ext))

    return sorted([archivo.name for archivo in archivos])
