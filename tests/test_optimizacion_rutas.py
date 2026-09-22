# -*- coding: utf-8 -*-
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from herramientas_geograficas import crear_capa_poligonos, cruzar_puntos_poligonos, unir_capas_poligonos
from optimizacion_rutas import (
    detectar_columnas_base,
    exportar_planificacion,
    extraer_forecast,
    planificar,
)


def base_puntos(titulares, suplentes=0):
    filas = []
    for indice, lon in enumerate(titulares):
        filas.append({
            "MT FINAL": "MT1", "SELECCION": "T", "LATITUD": 4.6,
            "LONGITUD": lon, "PDV": f"T{indice}", "RefID": f"T{indice}",
        })
    for indice in range(suplentes):
        filas.append({
            "MT FINAL": "MT1", "SELECCION": f"S{indice % 4 + 1}",
            "LATITUD": 4.6, "LONGITUD": titulares[indice % len(titulares)] + 0.0001,
            "PDV": f"S{indice}", "RefID": f"S{indice}",
        })
    return pd.DataFrame(filas)


class OptimizacionRutasTests(unittest.TestCase):
    def test_detecta_aliases_y_forecast(self):
        base = pd.DataFrame({
            "Ruta": ["M1"], "Tipo Punto": ["T"], "Lat": [4.6], "Lng": [-74.1],
            "Cliente": ["A"], "Codigo": [1],
        })
        columnas = detectar_columnas_base(base)
        self.assertEqual(columnas["mt"], "Ruta")
        self.assertEqual(columnas["longitud"], "Lng")
        forecast = extraer_forecast(pd.DataFrame({"MT FINAL": ["M1"], "1": [2], 2: [3], "nota": ["x"]}))
        self.assertEqual(forecast, {"M1": {1: 2, 2: 3}})

    def test_asigna_cupos_y_deja_excedentes_sin_dia(self):
        base = base_puntos([-74.00, -74.001, -74.002, -75.0, -75.001, -76.0])
        resultado = planificar(base, pd.DataFrame({"MT FINAL": ["MT1"], 1: [2], 2: [1]}))
        titulares = [p for p in resultado.puntos if p.tipo == "Titular"]
        self.assertEqual(sum(p.dia is not None for p in titulares), 3)
        self.assertEqual(sorted(sum(p.dia == dia for p in titulares) for dia in (1, 2)), [1, 2])
        self.assertEqual(resultado.modo, "solo-titulares")

    def test_asigna_tres_suplentes_por_titular(self):
        base = base_puntos([-74.0, -75.0], suplentes=6)
        resultado = planificar(base, pd.DataFrame({"MT FINAL": ["MT1"], 1: [1], 2: [1]}))
        suplentes = [p for p in resultado.puntos if p.tipo == "Suplente" and p.dia]
        self.assertEqual(len(suplentes), 6)
        self.assertEqual(sorted(sum(p.dia == dia for p in suplentes) for dia in (1, 2)), [3, 3])

    def test_usa_tolerancia_solo_para_formar_grupos_mas_compactos(self):
        coordenadas = [i * 0.0001 for i in range(17)] + [1 + i * 0.0001 for i in range(23)]
        resultado = planificar(
            base_puntos(coordenadas),
            pd.DataFrame({"MT FINAL": ["MT1"], 1: [20], 2: [20]}),
        )
        cantidades = [sum(p.dia == dia for p in resultado.puntos) for dia in (1, 2)]
        self.assertEqual(sorted(cantidades), [17, 23])

    def test_exporta_sobre_copia_y_conserva_otras_hojas(self):
        base = base_puntos([-74.0, -75.0])
        resultado = planificar(base, pd.DataFrame({"MT FINAL": ["MT1"], 1: [1]}))
        with tempfile.TemporaryDirectory() as temporal:
            origen = Path(temporal) / "base.xlsx"
            salida = Path(temporal) / "salida.xlsx"
            libro = Workbook()
            hoja = libro.active
            hoja.title = "Base"
            hoja.append(list(base.columns))
            for fila in base.itertuples(index=False, name=None):
                hoja.append(fila)
            control = libro.create_sheet("Control")
            control["A1"] = "CONSERVAR"
            libro.save(origen)
            exportar_planificacion(origen, salida, resultado)
            guardado = load_workbook(salida, data_only=False)
            self.assertEqual(guardado["Control"]["A1"].value, "CONSERVAR")
            encabezados = [celda.value for celda in guardado["Base"][1]]
            self.assertIn("DIA", encabezados)
            self.assertIn("Promedio metros", encabezados)

    def test_cruce_puntos_con_poligono_dibujado(self):
        puntos = pd.DataFrame({"LATITUD": [4.5, 7.0], "LONGITUD": [-74.0, -74.0]})
        capa = crear_capa_poligonos([{
            "id": 1, "descripcion": "Centro",
            "coordenadas": [(-75, 4), (-73, 4), (-73, 5), (-75, 5), (-75, 4)],
        }])
        salida = cruzar_puntos_poligonos(puntos, "LATITUD", "LONGITUD", capa)
        self.assertEqual(salida["DENTRO_POLIGONO"].tolist(), ["SI", "NO"])
        self.assertEqual(salida.loc[0, "DESCRIPCION"], "Centro")

    def test_capa_vacia_es_valida_antes_de_cargar_archivos(self):
        capa = unir_capas_poligonos(None, [])
        self.assertTrue(capa.empty)
        self.assertEqual(str(capa.crs), "EPSG:4326")


if __name__ == "__main__":
    unittest.main()
