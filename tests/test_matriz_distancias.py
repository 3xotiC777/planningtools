# -*- coding: utf-8 -*-
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from matriz_distancias import (
    NOMBRE_DISTANCIA,
    NOMBRE_RANGO,
    calcular_matriz_distancias,
    calcular_distancias_por_cercania,
    exportar_matriz_distancias,
    exportar_ruta_optima,
    ordenar_ruta_por_cercania,
)


class MatrizDistanciasTests(unittest.TestCase):
    def _archivos(self, carpeta: Path, origen: pd.DataFrame, destino: pd.DataFrame):
        ruta_o = carpeta / "origen.xlsx"
        ruta_d = carpeta / "destino.xlsx"
        origen.to_excel(ruta_o, index=False)
        destino.to_excel(ruta_d, index=False)
        return ruta_o, ruta_d

    def test_cruza_por_puente_y_calcula_siempre_en_metros(self):
        with tempfile.TemporaryDirectory() as temporal:
            origen, destino = self._archivos(
                Path(temporal),
                pd.DataFrame({"RefID": ["A", "B"], "LATITUD": [0.0, 1.0], "LONGITUD": [0.0, 0.0]}),
                pd.DataFrame({"RefID": ["A", "B"], "LATITUD": [0.0, 1.0], "LONGITUD": [1.0, 0.0]}),
            )
            resultado = calcular_matriz_distancias(origen, destino)
            self.assertEqual(resultado.shape[0], 2)
            self.assertTrue(111_000 <= float(resultado.loc[0, NOMBRE_DISTANCIA]) <= 112_000)
            self.assertEqual(float(resultado.loc[1, NOMBRE_DISTANCIA]), 0.0)
            self.assertEqual(resultado.loc[0, NOMBRE_RANGO], ">1000 m")

    def test_permite_elegir_puentes_gps_y_columnas_de_salida(self):
        with tempfile.TemporaryDirectory() as temporal:
            origen, destino = self._archivos(
                Path(temporal),
                pd.DataFrame({
                    "Código tienda": [101, 102], "Nombre": ["Uno", "Dos"],
                    "Y cliente": [9.0, 8.0], "X cliente": [-79.0, -80.0], "No incluir": [1, 2],
                }),
                pd.DataFrame({
                    "ID PDV": ["101.0", "102"], "País": ["Panamá", "Panamá"],
                    "Latitud export": [9.0001, 8.0], "Longitud export": [-79.0, -80.0], "Extra": [3, 4],
                }),
            )
            resultado = calcular_matriz_distancias(
                origen, destino,
                puente_origen="Código tienda", puente_destino="ID PDV",
                latitud_origen="Y cliente", longitud_origen="X cliente",
                latitud_destino="Latitud export", longitud_destino="Longitud export",
                columnas_origen=["Código tienda", "Nombre"],
                columnas_destino=["ID PDV", "País"],
            )
            self.assertEqual(
                list(resultado.columns),
                ["Código tienda", "Nombre", "ID PDV", "País", NOMBRE_DISTANCIA, NOMBRE_RANGO],
            )
            self.assertNotIn("No incluir", resultado.columns)
            self.assertNotIn("Extra", resultado.columns)
            self.assertLess(float(resultado.loc[0, NOMBRE_DISTANCIA]), 20)

    def test_el_cruce_es_lineal_para_llaves_unicas(self):
        with tempfile.TemporaryDirectory() as temporal:
            cantidad = 500
            origen, destino = self._archivos(
                Path(temporal),
                pd.DataFrame({"ID": range(cantidad), "LAT": [8.0] * cantidad, "LON": [-80.0] * cantidad}),
                pd.DataFrame({"ID": range(cantidad), "LAT": [8.0] * cantidad, "LON": [-80.0] * cantidad}),
            )
            resultado = calcular_matriz_distancias(origen, destino)
            self.assertEqual(len(resultado), cantidad)
            self.assertTrue((resultado[NOMBRE_DISTANCIA] == 0).all())

    def test_llaves_repetidas_se_emparejan_por_orden_sin_producto_cartesiano(self):
        with tempfile.TemporaryDirectory() as temporal:
            origen, destino = self._archivos(
                Path(temporal),
                pd.DataFrame({"ID": [1, 1], "LAT": [8.0, 9.0], "LON": [-80.0, -79.0]}),
                pd.DataFrame({"ID": [1, 1], "LAT": [8.0, 9.0], "LON": [-80.0, -79.0]}),
            )
            resultado = calcular_matriz_distancias(origen, destino)
            self.assertEqual(len(resultado), 2)
            self.assertTrue((resultado[NOMBRE_DISTANCIA] == 0).all())

    def test_exporta_hojas_resumen_y_bd_como_el_archivo_de_referencia(self):
        with tempfile.TemporaryDirectory() as temporal:
            carpeta = Path(temporal)
            origen, destino = self._archivos(
                carpeta,
                pd.DataFrame({"ID": [1, 2], "LAT": [8.0, 8.0], "LON": [-80.0, -80.0]}),
                pd.DataFrame({"ID": [1, 2], "LAT": [8.0, 8.01], "LON": [-80.0, -80.0]}),
            )
            resultado = calcular_matriz_distancias(origen, destino)
            salida = exportar_matriz_distancias(carpeta / "Comparativa_Distancias_GPS.xlsx", resultado)
            libro = load_workbook(salida, read_only=True, data_only=True)
            try:
                self.assertEqual(libro.sheetnames, ["RESUMEN", "BD"])
                filas = list(libro["BD"].iter_rows(values_only=True))
                encabezados = list(filas[0])
                self.assertIn(NOMBRE_DISTANCIA, encabezados)
                self.assertIn(NOMBRE_RANGO, encabezados)
                self.assertEqual(len(filas), 3)
            finally:
                libro.close()

    def test_cercania_no_necesita_columna_puente(self):
        with tempfile.TemporaryDirectory() as temporal:
            origen, destino = self._archivos(
                Path(temporal),
                pd.DataFrame({"Punto": ["A", "B"], "LAT": [0.0, 0.0], "LON": [0.1, 9.9]}),
                pd.DataFrame({"Sitio": ["Oeste", "Este"], "LAT": [0.0, 0.0], "LON": [0.0, 10.0]}),
            )
            resultado = calcular_distancias_por_cercania(
                origen, destino,
                columnas_analisis=["Punto", "LAT", "LON"],
                columnas_referencia=["Sitio"],
            )
            self.assertEqual(resultado["Sitio"].tolist(), ["Oeste", "Este"])
            self.assertTrue(resultado[NOMBRE_DISTANCIA].between(11_000, 11_200).all())

            solo_analisis = calcular_distancias_por_cercania(
                origen, destino,
                columnas_analisis=["Punto"], columnas_referencia=[],
            )
            self.assertEqual(list(solo_analisis.columns), ["Punto", NOMBRE_DISTANCIA, NOMBRE_RANGO])

    def test_ruta_se_reinicia_por_grupo_y_sigue_el_vecino_mas_cercano(self):
        with tempfile.TemporaryDirectory() as temporal:
            carpeta = Path(temporal)
            archivo = carpeta / "rutas.xlsx"
            pd.DataFrame({
                "ID": ["A", "B", "C", "D"],
                "COMUNA": ["Uno", "Uno", "Uno", "Dos"],
                "LAT": [0.0, 0.0, 0.0, 5.0],
                "LON": [0.0, 2.0, 1.0, 5.0],
            }).to_excel(archivo, index=False)
            resultado = ordenar_ruta_por_cercania(
                archivo, columna_reinicio="COMUNA",
                columna_latitud="LAT", columna_longitud="LON",
            )
            self.assertEqual(resultado["ID"].tolist(), ["D", "A", "C", "B"])
            self.assertEqual(resultado["Orden_Ruta"].tolist(), [1, 1, 2, 3])
            self.assertEqual(resultado["Distancia_Punto_Anterior_Metros"].iloc[0], 0)
            self.assertEqual(resultado["Distancia_Punto_Anterior_Metros"].iloc[1], 0)
            self.assertTrue(111_000 <= resultado["Distancia_Punto_Anterior_Metros"].iloc[2] <= 112_000)

            salida = exportar_ruta_optima(carpeta / "Distancias_Ruta_Optima.xlsx", resultado)
            libro = load_workbook(salida, read_only=True, data_only=True)
            try:
                self.assertEqual(libro.sheetnames, ["Datos_Ordenados_Ruta", "Total_Por_Comuna"])
            finally:
                libro.close()


if __name__ == "__main__":
    unittest.main()
