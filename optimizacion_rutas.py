# -*- coding: utf-8 -*-
"""Motor de planificación geográfica portado de Control_Dispersion.

El módulo no depende de la interfaz. Lee las dos bases de Excel, detecta sus
columnas, forma jornadas compactas respetando el forecast, asigna suplentes y
permite escribir el resultado sobre una copia del archivo original.
"""

from __future__ import annotations

import math
import json
import re
import time
import unicodedata
from dataclasses import dataclass, replace
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Callable
from urllib.parse import quote
from urllib.request import urlopen

import pandas as pd
from openpyxl import load_workbook


MAX_DISTANCIA_SUPLENTE_METROS = 15_000
TOLERANCIA_FORECAST_MAXIMA = 5


class ErrorOptimizacion(ValueError):
    """Error de datos comprensible para el usuario final."""


@dataclass
class PuntoRuta:
    id: str
    indice_origen: int
    ref_id: str
    nombre: str
    mt: str
    seleccion: str
    tipo: str
    prioridad: int
    lat: float
    lon: float
    dia: int | None = None
    mt_asignado: str | None = None
    promedio_metros: float | None = None


@dataclass
class ResultadoPlanificacion:
    puntos: list[PuntoRuta]
    forecast: dict[str, dict[int, int]]
    modo: str
    avisos: list[str]
    filas_origen: pd.DataFrame
    columnas_base: dict[str, str]


def normalizar(valor: Any) -> str:
    texto = unicodedata.normalize("NFD", str("" if valor is None else valor))
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", texto.strip().upper())


def llave(valor: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", normalizar(valor))


def _numero(valor: Any) -> float | None:
    if valor is None or (isinstance(valor, float) and math.isnan(valor)):
        return None
    try:
        numero = float(str(valor).strip().replace(",", "."))
    except (TypeError, ValueError):
        return None
    return numero if math.isfinite(numero) else None


def _columna(columnas: list[str], nombres: list[str], contiene: tuple[str, ...] = ()) -> str | None:
    buscadas = {llave(nombre) for nombre in nombres}
    exacta = next((col for col in columnas if llave(col) in buscadas), None)
    if exacta:
        return exacta
    return next((col for col in columnas if any(fragmento in llave(col) for fragmento in contiene)), None)


def detectar_columnas_base(datos: pd.DataFrame) -> dict[str, str]:
    if datos.empty:
        raise ErrorOptimizacion("La base de puntos no contiene registros.")
    columnas = [str(col) for col in datos.columns]
    encontradas = {
        "mt": _columna(columnas, ["MT FINAL", "MT", "TERRITORIO", "ZONA", "RUTA"], ("MT", "TERRITORIO")),
        "seleccion": _columna(columnas, ["SELECCION", "SELECCION PUNTO", "TIPO", "TIPO PUNTO", "KIND", "CLASE", "CATEGORIA"], ("SELECC", "TIPO")),
        "latitud": _columna(columnas, ["LATITUDE", "LATITUD", "LAT", "Y"], ("LAT",)),
        "longitud": _columna(columnas, ["LONGITUDE", "LONGITUD", "LON", "LNG", "X"], ("LON", "LNG")),
        "nombre": _columna(columnas, ["PDV", "NOMBRE", "CLIENTE", "PUNTO", "NAME", "DESCRIPCION", "ESTABLECIMIENTO"], ("PDV", "NOMBRE", "CLIENTE")),
        "ref_id": _columna(columnas, ["REFID", "ID", "CODIGO", "COD", "REF", "PUNTO ID"], ("REF", "COD")),
    }
    etiquetas = {
        "mt": "MT FINAL", "seleccion": "SELECCION", "latitud": "LATITUD",
        "longitud": "LONGITUD", "nombre": "PDV", "ref_id": "RefID",
    }
    faltante = next((campo for campo, col in encontradas.items() if not col), None)
    if faltante:
        raise ErrorOptimizacion(f'No se encontró la columna "{etiquetas[faltante]}" en la base de puntos.')
    return {campo: str(col) for campo, col in encontradas.items()}


def detectar_columnas_coordenadas(datos: pd.DataFrame) -> tuple[str, str]:
    """Detección más permisiva para la opción Puntos en Excel."""
    columnas = [str(col) for col in datos.columns]
    lat = _columna(columnas, ["LATITUDE", "LATITUD", "LAT", "Y"], ("LAT",))
    lon = _columna(columnas, ["LONGITUDE", "LONGITUD", "LON", "LNG", "X"], ("LON", "LNG"))
    if not lat or not lon:
        raise ErrorOptimizacion("El archivo debe tener columnas LATITUD y LONGITUD (también se aceptan LAT/LON o Y/X).")
    return lat, lon


def _prioridad_suplente(seleccion: str, pais: str) -> int:
    orden = ["S PANEL", "S1", "S2", "S3", "S4", "S ON", "S ORO"] if "DOMINIC" in pais else ["S1", "S2", "S3", "S4"]
    try:
        return orden.index(seleccion)
    except ValueError:
        return 99


def extraer_puntos(datos: pd.DataFrame) -> tuple[list[PuntoRuta], dict[str, str], str]:
    campos = detectar_columnas_base(datos)
    columnas = [str(col) for col in datos.columns]
    col_pais = _columna(columnas, ["PAIS", "COUNTRY"])
    coordenadas: list[tuple[float | None, float | None]] = []
    validas = intercambiadas = 0
    for _, fila in datos.iterrows():
        lat = _numero(fila[campos["latitud"]])
        lon = _numero(fila[campos["longitud"]])
        coordenadas.append((lat, lon))
        if lat is not None and lon is not None:
            validas += 1
            if abs(lat) > 60 and abs(lon) < 60:
                intercambiadas += 1
    invertir = validas > 0 and intercambiadas > validas * 0.55

    puntos: list[PuntoRuta] = []
    hay_suplentes = False
    for posicion, (_, fila) in enumerate(datos.iterrows()):
        lat, lon = coordenadas[posicion]
        if lat is None or lon is None:
            continue
        if invertir:
            lat, lon = lon, lat
        if not (-90 <= lat <= 90 and -180 <= lon <= 180):
            continue
        seleccion = normalizar(fila[campos["seleccion"]])
        if seleccion in {"T", "T PANEL"} or "TITULAR" in seleccion:
            tipo = "Titular"
        elif seleccion.startswith("S") or "SUPLENTE" in seleccion:
            tipo = "Suplente"
            hay_suplentes = True
        else:
            continue
        ref = str(fila[campos["ref_id"]] if pd.notna(fila[campos["ref_id"]]) else posicion + 1)
        nombre = str(fila[campos["nombre"]] if pd.notna(fila[campos["nombre"]]) else ref)
        pais = normalizar(fila[col_pais]) if col_pais else ""
        puntos.append(PuntoRuta(
            id=f"{ref}-{posicion}", indice_origen=posicion, ref_id=ref, nombre=nombre,
            mt=str(fila[campos["mt"]]).strip(), seleccion=seleccion, tipo=tipo,
            prioridad=_prioridad_suplente(seleccion, pais), lat=float(lat), lon=float(lon),
        ))
    if not puntos:
        raise ErrorOptimizacion("No se encontraron titulares o suplentes con coordenadas válidas.")
    return puntos, campos, "con-suplentes" if hay_suplentes else "solo-titulares"


def extraer_forecast(datos: pd.DataFrame) -> dict[str, dict[int, int]]:
    if datos.empty:
        raise ErrorOptimizacion("El forecast no contiene registros.")
    columnas = [str(col) for col in datos.columns]
    col_mt = _columna(columnas, ["MT FINAL"])
    if not col_mt:
        raise ErrorOptimizacion('No se encontró la columna "MT FINAL" en el forecast.')
    dias: list[tuple[Any, int]] = []
    for columna in datos.columns:
        try:
            dia = int(float(str(columna)))
        except (TypeError, ValueError):
            continue
        if dia > 0:
            dias.append((columna, dia))
    if not dias:
        raise ErrorOptimizacion("No se encontraron columnas de días numéricos en el forecast.")
    resultado: dict[str, dict[int, int]] = {}
    for _, fila in datos.iterrows():
        mt = str(fila[col_mt]).strip()
        if not mt or mt.lower() == "nan":
            continue
        for columna, dia in dias:
            cantidad = _numero(fila[columna])
            if cantidad is not None and cantidad > 0:
                resultado.setdefault(mt, {})[dia] = int(round(cantidad))
    if not resultado:
        raise ErrorOptimizacion("El forecast no contiene cantidades positivas por MT FINAL y día.")
    return resultado


def distancia_metros(a: PuntoRuta, b: PuntoRuta) -> float:
    radio = 6_371_000.0
    lat1, lat2 = math.radians(a.lat), math.radians(b.lat)
    dlat = math.radians(b.lat - a.lat)
    dlon = math.radians(b.lon - a.lon)
    h = math.sin(dlat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon / 2) ** 2
    return 2 * radio * math.asin(math.sqrt(min(1.0, h)))


def tolerancia_forecast(esperado: int) -> int:
    return min(TOLERANCIA_FORECAST_MAXIMA, max(1, int(esperado * 0.25)))


def _matriz(puntos: list[PuntoRuta]) -> list[list[float]]:
    return [[distancia_metros(a, b) for b in puntos] for a in puntos]


def _subconjunto_denso(puntos: list[PuntoRuta], objetivo: int) -> list[PuntoRuta]:
    if len(puntos) <= objetivo:
        return list(puntos)
    matriz = _matriz(puntos)
    vecinos = min(6, len(puntos) - 1)
    puntajes = []
    for indice, punto in enumerate(puntos):
        cercanos = sorted(valor for otro, valor in enumerate(matriz[indice]) if otro != indice)[:vecinos]
        puntajes.append((sum(cercanos), indice, punto))
    return [punto for _, _, punto in sorted(puntajes)[:objetivo]]


def _medoides_iniciales(matriz: list[list[float]], cantidad: int) -> list[int]:
    vecinos = min(5, len(matriz) - 1)
    densidad = [sum(sorted(v for j, v in enumerate(fila) if j != i)[:vecinos]) for i, fila in enumerate(matriz)]
    medoides = [min(range(len(matriz)), key=lambda i: densidad[i])]
    while len(medoides) < cantidad:
        candidatos = [i for i in range(len(matriz)) if i not in medoides]
        medoides.append(max(candidatos, key=lambda i: min(matriz[i][m] for m in medoides)))
    return medoides


def _asignar_capacidad(matriz: list[list[float]], medoides: list[int], capacidades: list[int]) -> list[int]:
    etiquetas = [-1] * len(matriz)
    restantes = list(capacidades)
    pendientes = set(range(len(matriz)))
    while pendientes:
        mejor: tuple[float, float, int, int] | None = None
        for punto in pendientes:
            costos = sorted((matriz[punto][med], cluster) for cluster, med in enumerate(medoides) if restantes[cluster] > 0)
            if not costos:
                continue
            costo, cluster = costos[0]
            arrepentimiento = (costos[1][0] - costo) if len(costos) > 1 else float("inf")
            candidato = (arrepentimiento, -costo, -punto, cluster)
            if mejor is None or candidato > mejor:
                mejor = candidato
        if mejor is None:
            break
        punto = -mejor[2]
        cluster = mejor[3]
        etiquetas[punto] = cluster
        restantes[cluster] -= 1
        pendientes.remove(punto)
    return etiquetas


def _refinar_intercambios(etiquetas: list[int], matriz: list[list[float]], maximos: int = 90) -> int:
    intercambios = 0
    while intercambios < maximos:
        mejor_delta = -0.001
        mejor_par: tuple[int, int] | None = None
        for a in range(len(etiquetas)):
            for b in range(a + 1, len(etiquetas)):
                ca, cb = etiquetas[a], etiquetas[b]
                if ca == cb:
                    continue
                delta = 0.0
                for otro, cluster in enumerate(etiquetas):
                    if otro in (a, b):
                        continue
                    if cluster == ca:
                        delta += matriz[b][otro] - matriz[a][otro]
                    elif cluster == cb:
                        delta += matriz[a][otro] - matriz[b][otro]
                if delta < mejor_delta:
                    mejor_delta, mejor_par = delta, (a, b)
        if mejor_par is None:
            break
        a, b = mejor_par
        etiquetas[a], etiquetas[b] = etiquetas[b], etiquetas[a]
        intercambios += 1
    return intercambios


def _costo_grupo(miembros: list[int], matriz_cuadrada: list[list[float]]) -> float:
    if len(miembros) < 2:
        return 0.0
    valores = [matriz_cuadrada[a][b] for pos, a in enumerate(miembros) for b in miembros[pos + 1:]]
    return sum(valores) / len(valores) + max(valores)


def _refinar_con_tolerancia(
    etiquetas: list[int],
    matriz: list[list[float]],
    esperados: list[int],
) -> int:
    """Usa hasta ±5 puestos solamente cuando baja la dispersión."""
    cuadrados = [[valor * valor for valor in fila] for fila in matriz]
    inferiores = [max(1, valor - tolerancia_forecast(valor)) for valor in esperados]
    superiores = [valor + tolerancia_forecast(valor) for valor in esperados]
    movidos: set[int] = set()
    while True:
        miembros = [[i for i, etiqueta in enumerate(etiquetas) if etiqueta == grupo] for grupo in range(len(esperados))]
        mejor: tuple[float, int, int] | None = None
        for origen in range(len(esperados)):
            if len(miembros[origen]) <= inferiores[origen]:
                continue
            for destino in range(len(esperados)):
                if origen == destino or len(miembros[destino]) >= superiores[destino]:
                    continue
                antes = _costo_grupo(miembros[origen], cuadrados) + _costo_grupo(miembros[destino], cuadrados)
                for punto in miembros[origen]:
                    if punto in movidos:
                        continue
                    origen_nuevo = [i for i in miembros[origen] if i != punto]
                    destino_nuevo = miembros[destino] + [punto]
                    despues = _costo_grupo(origen_nuevo, cuadrados) + _costo_grupo(destino_nuevo, cuadrados)
                    delta = despues - antes
                    if delta < -1 and (mejor is None or delta < mejor[0]):
                        mejor = (delta, punto, destino)
        if mejor is None:
            break
        _, punto, destino = mejor
        etiquetas[punto] = destino
        movidos.add(punto)
    return len(movidos)


def _clusters_capacitados(puntos: list[PuntoRuta], capacidades: list[int]) -> list[int]:
    if not puntos:
        return []
    if len(capacidades) == 1:
        return [0] * len(puntos)
    matriz = _matriz(puntos)
    medoides = _medoides_iniciales(matriz, len(capacidades))
    etiquetas = _asignar_capacidad(matriz, medoides, capacidades)
    for _ in range(10):
        siguientes: list[int] = []
        for cluster, anterior in enumerate(medoides):
            miembros = [i for i, etiqueta in enumerate(etiquetas) if etiqueta == cluster]
            siguientes.append(min(miembros, key=lambda i: sum(matriz[i][j] for j in miembros)) if miembros else anterior)
        if siguientes == medoides:
            break
        medoides = siguientes
        etiquetas = _asignar_capacidad(matriz, medoides, capacidades)
    cuadrados = [[valor * valor for valor in fila] for fila in matriz]
    _refinar_intercambios(etiquetas, matriz, 36 if len(puntos) > 300 else 90)
    _refinar_intercambios(etiquetas, cuadrados, 48 if len(puntos) > 300 else 120)
    _refinar_con_tolerancia(etiquetas, matriz, capacidades)
    return etiquetas


def _centro_grupo(grupo: list[PuntoRuta]) -> PuntoRuta:
    return min(grupo, key=lambda p: sum(distancia_metros(p, q) for q in grupo))


def _secuenciar_dias(puntos: list[PuntoRuta], forecast: dict[str, dict[int, int]]) -> int:
    """Renombra grupos completos siguiendo una ruta geográfica cercana."""
    cambios = 0
    mts = sorted({p.mt_asignado for p in puntos if p.tipo == "Titular" and p.dia and p.mt_asignado})
    for mt in mts:
        diarios: dict[int, list[PuntoRuta]] = {}
        for punto in puntos:
            if punto.tipo == "Titular" and punto.mt_asignado == mt and punto.dia:
                diarios.setdefault(punto.dia, []).append(punto)
        dias_destino = sorted(diarios)
        if len(dias_destino) < 2:
            continue
        pendientes = [(dia, grupo, _centro_grupo(grupo)) for dia, grupo in diarios.items()]
        # Comienza por el grupo más occidental y continúa por proximidad; no
        # parte jornadas y conserva la cuota compatible del forecast.
        actual = min(pendientes, key=lambda x: (x[2].lon, x[2].lat))
        orden = [actual]
        pendientes.remove(actual)
        while pendientes:
            posicion = len(orden)
            dia_objetivo = dias_destino[posicion]
            esperado = forecast.get(str(mt), {}).get(dia_objetivo, len(pendientes[0][1]))
            compatibles = [g for g in pendientes if abs(len(g[1]) - esperado) <= tolerancia_forecast(esperado)]
            candidatos = compatibles or pendientes
            actual = min(candidatos, key=lambda x: distancia_metros(orden[-1][2], x[2]))
            orden.append(actual)
            pendientes.remove(actual)
        for dia_nuevo, (dia_anterior, grupo, _) in zip(dias_destino, orden):
            if dia_nuevo != dia_anterior:
                cambios += 1
            for punto in grupo:
                punto.dia = dia_nuevo
    return cambios


def _asignar_suplentes(puntos: list[PuntoRuta], avisos: list[str]) -> None:
    grupos: list[dict[str, Any]] = []
    claves = sorted({(p.mt_asignado, p.dia) for p in puntos if p.tipo == "Titular" and p.mt_asignado and p.dia})
    for mt, dia in claves:
        titulares = [p for p in puntos if p.tipo == "Titular" and p.mt_asignado == mt and p.dia == dia]
        grupos.append({"mt": mt, "dia": dia, "titulares": titulares, "cupo": len(titulares) * 3, "asignados": [], "tiene_cercanos": False})
    suplentes = sorted((p for p in puntos if p.tipo == "Suplente"), key=lambda p: (p.prioridad, p.id))
    disponibles = {p.id: p for p in suplentes}
    aristas: list[tuple[int, float, str, int]] = []
    for indice, grupo in enumerate(grupos):
        for suplente in suplentes:
            distancia = min(distancia_metros(suplente, titular) for titular in grupo["titulares"])
            if distancia <= MAX_DISTANCIA_SUPLENTE_METROS:
                grupo["tiene_cercanos"] = True
                aristas.append((suplente.prioridad, distancia, suplente.id, indice))
    for _, _, punto_id, indice in sorted(aristas):
        grupo = grupos[indice]
        if punto_id not in disponibles or len(grupo["asignados"]) >= grupo["cupo"]:
            continue
        punto = disponibles.pop(punto_id)
        punto.mt_asignado, punto.dia = grupo["mt"], grupo["dia"]
        grupo["asignados"].append(punto)
    for grupo in grupos:
        faltan = grupo["cupo"] - len(grupo["asignados"])
        if faltan > 0 and disponibles and not grupo["tiene_cercanos"]:
            # Fallback fiel al origen: solo usa lejanos cuando no existe ningún
            # candidato cercano para completar esa jornada.
            cercanos = sorted(
                disponibles.values(),
                key=lambda p: (min(distancia_metros(p, t) for t in grupo["titulares"]), p.prioridad, p.id),
            )[:faltan]
            for punto in cercanos:
                disponibles.pop(punto.id, None)
                punto.mt_asignado, punto.dia = grupo["mt"], grupo["dia"]
                grupo["asignados"].append(punto)
            if cercanos:
                avisos.append(f'{grupo["mt"]}, día {grupo["dia"]}: se usaron {len(cercanos)} suplentes fuera de 15 km.')
        if len(grupo["asignados"]) < grupo["cupo"]:
            avisos.append(f'{grupo["mt"]}, día {grupo["dia"]}: {len(grupo["asignados"])}/{grupo["cupo"]} suplentes disponibles.')


def recalcular_promedios(puntos: list[PuntoRuta]) -> None:
    for punto in puntos:
        punto.promedio_metros = None
    grupos = {(p.mt_asignado, p.dia) for p in puntos if p.mt_asignado and p.dia}
    for mt, dia in grupos:
        miembros = [p for p in puntos if p.mt_asignado == mt and p.dia == dia]
        titulares = [p for p in miembros if p.tipo == "Titular"]
        for punto in miembros:
            if not titulares:
                continue
            if punto.tipo == "Titular":
                otros = [t for t in titulares if t.id != punto.id]
                punto.promedio_metros = sum(distancia_metros(punto, t) for t in otros) / len(otros) if otros else 0.0
            else:
                punto.promedio_metros = min(distancia_metros(punto, t) for t in titulares)


def planificar(
    base: pd.DataFrame,
    forecast_df: pd.DataFrame,
    progreso: Callable[[float, str], None] | None = None,
    qa_vial_automatico: bool = False,
) -> ResultadoPlanificacion:
    avisar = progreso or (lambda _p, _m: None)
    avisar(0.05, "Validando columnas y coordenadas...")
    puntos, columnas, modo = extraer_puntos(base)
    forecast = extraer_forecast(forecast_df)
    siguientes = [replace(p, dia=None, mt_asignado=None, promedio_metros=None) for p in puntos]
    avisos: list[str] = []
    por_mt: dict[str, list[PuntoRuta]] = {}
    for punto in siguientes:
        por_mt.setdefault(punto.mt, []).append(punto)
    total_mt = max(1, len(forecast))
    for numero, (mt, diario) in enumerate(forecast.items(), start=1):
        avisar(0.1 + 0.65 * numero / total_mt, f"Formando jornadas compactas de {mt}...")
        disponibles = por_mt.get(mt, [])
        titulares = [p for p in disponibles if p.tipo == "Titular"]
        dias = sorted((int(dia), int(cantidad)) for dia, cantidad in diario.items() if cantidad > 0)
        necesarios = sum(cantidad for _, cantidad in dias)
        if not disponibles:
            avisos.append(f"{mt}: no hay puntos con coordenadas en la base.")
            continue
        if len(titulares) < necesarios:
            avisos.append(f"{mt}: el forecast pide {necesarios} titulares y la base tiene {len(titulares)}.")
        elegidos = _subconjunto_denso(titulares, min(necesarios, len(titulares)))
        restantes = len(elegidos)
        efectivos: list[tuple[int, int]] = []
        for dia, cantidad in dias:
            cupo = min(cantidad, restantes)
            restantes -= cupo
            if cupo:
                efectivos.append((dia, cupo))
        etiquetas = _clusters_capacitados(elegidos, [cupo for _, cupo in efectivos]) if efectivos else []
        for indice, punto in enumerate(elegidos):
            cluster = etiquetas[indice]
            if cluster >= 0:
                punto.dia, punto.mt_asignado = efectivos[cluster][0], mt
    if qa_vial_automatico:
        avisar(0.74, "Detectando cruces que necesitan QA vial...")
        avisos.extend(qa_vial_inteligente(siguientes, progreso=lambda texto: avisar(0.76, texto)))
    cambios = _secuenciar_dias(siguientes, forecast)
    if cambios:
        avisos.insert(0, f"Secuencia geográfica final: se renumeraron {cambios} grupos completos sin separar sus puntos.")
    if modo == "con-suplentes":
        avisar(0.82, "Asignando suplentes cercanos (relación 1:3)...")
        _asignar_suplentes(siguientes, avisos)
    else:
        avisos.insert(0, "Modo solo titulares: se usó el forecast por MT FINAL y día; los excedentes quedan sin día.")
    recalcular_promedios(siguientes)
    avisar(1.0, "Planificación terminada.")
    return ResultadoPlanificacion(siguientes, forecast, modo, avisos, base.copy(), columnas)


def planificar_archivos(
    ruta_base: str | Path,
    ruta_forecast: str | Path,
    progreso: Callable[[float, str], None] | None = None,
    qa_vial_automatico: bool = True,
) -> ResultadoPlanificacion:
    try:
        base = pd.read_excel(ruta_base)
    except Exception as exc:
        raise ErrorOptimizacion(f"No fue posible leer la base de puntos: {exc}") from exc
    try:
        forecast = pd.read_excel(ruta_forecast)
    except Exception as exc:
        raise ErrorOptimizacion(f"No fue posible leer el forecast: {exc}") from exc
    return planificar(base, forecast, progreso, qa_vial_automatico=qa_vial_automatico)


def mover_puntos(resultado: ResultadoPlanificacion, ids: set[str], dia: int, mt: str | None = None) -> None:
    for punto in resultado.puntos:
        if punto.id in ids:
            punto.dia = int(dia)
            punto.mt_asignado = mt or punto.mt_asignado or punto.mt
    recalcular_promedios(resultado.puntos)


def exportar_planificacion(
    ruta_base: str | Path,
    ruta_salida: str | Path,
    resultado: ResultadoPlanificacion,
) -> Path:
    """Conserva hojas y formato del libro base y agrega los resultados."""
    libro = load_workbook(ruta_base)
    hoja = libro[libro.sheetnames[0]]
    encabezados = {llave(celda.value): celda.column for celda in hoja[1] if celda.value is not None}

    def asegurar(nombre: str) -> int:
        clave = llave(nombre)
        if clave in encabezados:
            return encabezados[clave]
        columna = hoja.max_column + 1
        hoja.cell(1, columna, nombre)
        encabezados[clave] = columna
        return columna

    col_dia = asegurar("DIA")
    col_promedio = asegurar("Promedio metros")
    col_mt = asegurar("MT FINAL")
    por_indice = {p.indice_origen: p for p in resultado.puntos}
    for indice in range(len(resultado.filas_origen)):
        punto = por_indice.get(indice)
        fila = indice + 2
        if punto is None:
            hoja.cell(fila, col_dia, None)
            hoja.cell(fila, col_promedio, None)
            continue
        hoja.cell(fila, col_dia, punto.dia)
        hoja.cell(fila, col_promedio, round(punto.promedio_metros or 0.0, 1) if punto.dia else None)
        if punto.mt_asignado:
            hoja.cell(fila, col_mt, punto.mt_asignado)
    salida = Path(ruta_salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    libro.save(salida)
    return salida


def promedio_vecino_mas_cercano(datos: pd.DataFrame, col_lat: str, col_lon: str) -> float | None:
    coordenadas = []
    for lat, lon in zip(pd.to_numeric(datos[col_lat], errors="coerce"), pd.to_numeric(datos[col_lon], errors="coerce")):
        if pd.notna(lat) and pd.notna(lon) and -90 <= lat <= 90 and -180 <= lon <= 180:
            coordenadas.append(PuntoRuta("", 0, "", "", "", "", "", 0, float(lat), float(lon)))
    if len(coordenadas) < 2:
        return None
    distancias = [min(distancia_metros(p, q) for j, q in enumerate(coordenadas) if i != j) for i, p in enumerate(coordenadas)]
    return sum(distancias) / len(distancias)


def _media_interna(etiquetas: list[int], matriz: list[list[float]]) -> float:
    valores = [matriz[a][b] for a in range(len(etiquetas)) for b in range(a + 1, len(etiquetas)) if etiquetas[a] == etiquetas[b]]
    return sum(valores) / len(valores) if valores else 0.0


def matriz_tiempos_viales(puntos: list[PuntoRuta]) -> list[list[float]]:
    """Consulta la API pública OSRM; las coordenadas se envían por HTTPS."""
    if len(puntos) > 225:
        raise ErrorOptimizacion("El QA vial admite hasta 225 titulares por MT en una ejecución.")
    resultado = [[0.0] * len(puntos) for _ in puntos]

    def consultar(origenes: list[PuntoRuta], destinos: list[PuntoRuta], offset_o: int, offset_d: int, mismos: bool) -> None:
        coordenadas = origenes if mismos else origenes + destinos
        texto = ";".join(f"{p.lon},{p.lat}" for p in coordenadas)
        if mismos:
            consulta = "annotations=duration&skip_waypoints=true"
        else:
            sources = ";".join(str(i) for i in range(len(origenes)))
            destinations = ";".join(str(i + len(origenes)) for i in range(len(destinos)))
            consulta = f"sources={sources}&destinations={destinations}&annotations=duration&skip_waypoints=true"
        url = f"https://router.project-osrm.org/table/v1/driving/{quote(texto, safe=',;-')}?{consulta}"
        try:
            with urlopen(url, timeout=12) as respuesta:  # nosec B310 - endpoint HTTPS fijo
                datos = json.loads(respuesta.read().decode("utf-8"))
        except Exception as exc:
            raise ErrorOptimizacion(f"No fue posible consultar la red vial de OSRM: {exc}") from exc
        if datos.get("code") != "Ok" or not datos.get("durations"):
            raise ErrorOptimizacion("El servicio vial no pudo construir la matriz de tiempos.")
        for fila, valores in enumerate(datos["durations"]):
            for columna, valor in enumerate(valores):
                if valor is None:
                    raise ErrorOptimizacion("Hay puntos que no pudieron conectarse a la red vial.")
                resultado[offset_o + fila][offset_d + columna] = float(valor)

    if len(puntos) <= 95:
        consultar(puntos, puntos, 0, 0, True)
    else:
        bloque = 45
        for inicio_o in range(0, len(puntos), bloque):
            for inicio_d in range(0, len(puntos), bloque):
                consultar(puntos[inicio_o:inicio_o + bloque], puntos[inicio_d:inicio_d + bloque], inicio_o, inicio_d, False)
                time.sleep(0.25)
    return [[(valor + resultado[b][a]) / 2 for b, valor in enumerate(fila)] for a, fila in enumerate(resultado)]


def _promedio_a(punto: PuntoRuta, grupo: list[PuntoRuta], excluir: bool = False) -> float:
    valores = [distancia_metros(punto, otro) for otro in grupo if not excluir or otro.id != punto.id]
    return sum(valores) / len(valores) if valores else 0.0


def _dispersion(grupo: list[PuntoRuta]) -> float:
    valores = [distancia_metros(a, b) for i, a in enumerate(grupo) for b in grupo[i + 1:]]
    return sum(valores) / len(valores) if valores else 0.0


def _candidatos_qa(puntos: list[PuntoRuta]) -> list[tuple[float, str, int, int]]:
    por_mt: dict[str, dict[int, list[PuntoRuta]]] = {}
    for punto in puntos:
        if punto.tipo == "Titular" and punto.dia and punto.mt_asignado:
            por_mt.setdefault(punto.mt_asignado, {}).setdefault(punto.dia, []).append(punto)
    candidatos = []
    for mt, diarios in por_mt.items():
        del_mt = []
        grupos = sorted(diarios.items())
        for i, (dia_a, a) in enumerate(grupos):
            for dia_b, b in grupos[i + 1:]:
                if len(a) + len(b) > 90:
                    continue
                dentro_a, dentro_b = _dispersion(a), _dispersion(b)
                ganancia_a = max((_promedio_a(p, a, True) - _promedio_a(p, b)) / max(_promedio_a(p, a, True), 1) for p in a)
                ganancia_b = max((_promedio_a(p, b, True) - _promedio_a(p, a)) / max(_promedio_a(p, b, True), 1) for p in b)
                cruce = min(distancia_metros(pa, pb) for pa in a for pb in b)
                dispersion = max(dentro_a, dentro_b, 1)
                severidad = 10 + ganancia_a + ganancia_b if min(ganancia_a, ganancia_b) > 0.04 else (1 - cruce / dispersion if dispersion > 2500 and cruce < dispersion * 0.3 else 0)
                if severidad > 0.08:
                    del_mt.append((severidad, mt, dia_a, dia_b))
        candidatos.extend(sorted(del_mt, reverse=True)[:2])
    return sorted(candidatos, reverse=True)[:24]


def qa_vial_inteligente(
    puntos: list[PuntoRuta],
    progreso: Callable[[str], None] | None = None,
) -> list[str]:
    """Revisa por carretera solo pares de jornadas espacialmente sospechosos."""
    candidatos = _candidatos_qa(puntos)
    if not candidatos:
        return ["QA vial automático: no se detectaron cruces espaciales que necesitaran validación."]
    revisados = fallidos = intercambios = cambiados = 0
    for numero, (_, mt, dia_a, dia_b) in enumerate(candidatos, start=1):
        if progreso:
            progreso(f"QA vial {numero}/{len(candidatos)}: {mt}, días {dia_a} y {dia_b}...")
        titulares = [p for p in puntos if p.tipo == "Titular" and p.mt_asignado == mt and p.dia in {dia_a, dia_b}]
        etiquetas = [0 if p.dia == dia_a else 1 for p in titulares]
        originales = list(etiquetas)
        try:
            matriz = matriz_tiempos_viales(titulares)
            antes = _media_interna(etiquetas, matriz)
            swaps = _refinar_intercambios(etiquetas, matriz, min(24, len(titulares)))
            despues = _media_interna(etiquetas, matriz)
            mejora = antes - despues
            if swaps and mejora >= 30 and mejora / max(antes, 1) >= 0.02:
                for punto, etiqueta in zip(titulares, etiquetas):
                    punto.dia = dia_a if etiqueta == 0 else dia_b
                intercambios += swaps
                cambiados += sum(a != b for a, b in zip(originales, etiquetas))
            revisados += 1
        except Exception:
            # Si el servicio no está disponible, insistir con todos los pares
            # solo alargaría el proceso sin cambiar la planificación.
            fallidos += len(candidatos) - numero + 1
            break
        if numero < len(candidatos):
            time.sleep(0.25)
    texto = (
        f"QA vial automático: {revisados} cruces revisados; {cambiados} titulares cambiaron de día mediante {intercambios} intercambios."
        if cambiados else
        f"QA vial automático: {revisados} cruces sospechosos revisados; la asignación ya era estable."
    )
    if fallidos:
        texto += f" {fallidos} validaciones no pudieron consultarse y conservaron su asignación."
    return [texto]


def ejecutar_qa_vial_mt(resultado: ResultadoPlanificacion, mt: str) -> str:
    titulares = [p for p in resultado.puntos if p.tipo == "Titular" and p.mt_asignado == mt and p.dia]
    if len(titulares) < 2:
        raise ErrorOptimizacion(f"{mt}: no hay suficientes titulares asignados para ejecutar el QA vial.")
    matriz = matriz_tiempos_viales(titulares)
    dias = sorted({p.dia for p in titulares if p.dia})
    indices = {dia: i for i, dia in enumerate(dias)}
    etiquetas = [indices[p.dia] for p in titulares]
    originales = list(etiquetas)
    antes = _media_interna(etiquetas, matriz)
    swaps = _refinar_intercambios(etiquetas, matriz, min(160, len(titulares) * 2))
    despues = _media_interna(etiquetas, matriz)
    for punto, etiqueta in zip(titulares, etiquetas):
        punto.dia = dias[etiqueta]
    recalcular_promedios(resultado.puntos)
    cambiados = sum(a != b for a, b in zip(originales, etiquetas))
    return (
        f"QA vial de {mt}: {cambiados} titulares cambiaron mediante {swaps} intercambios; "
        f"el tiempo interno bajó de {antes / 60:.1f} a {despues / 60:.1f} minutos."
        if swaps else f"QA vial de {mt}: la distribución ya era estable según los tiempos de conducción."
    )
